"""Synthetic inputs, an isolated SQLite deployment, and SQL extracts for the tutorial.

The notebook contains the model fitting and monitoring calls. This module never
connects to SQL Server or changes a database outside its marked demo directory.
"""

from __future__ import annotations

import json
from pathlib import Path
from tempfile import mkdtemp

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import text

_MARKER = ".sql-monitoring-demo"


def create_demo_directory(base: Path) -> Path:
    """Use a new directory on every run, preserving earlier demo output."""
    base = base.resolve()
    if base.exists():
        directory = Path(mkdtemp(prefix="run-", dir=base))
    else:
        base.mkdir(parents=True)
        directory = base
    (directory / _MARKER).write_text("Isolated synthetic SQL monitoring tutorial.\n")
    return directory


def synthetic_burn_cost(
    *, rows: int, seed: int, as_of: str, monitoring: bool = False
) -> pd.DataFrame:
    """Draw compound Poisson-gamma losses with Tweedie power 1.5.

    The target is loss per exposure. Exposure is the fitting weight. The later
    snapshot has modest claims inflation and a larger North region share.
    """
    rng = np.random.default_rng(seed)
    regions = ["North", "South", "East", "West"]
    region = rng.choice(regions, rows, p=[0.40, 0.20, 0.20, 0.20] if monitoring else [0.25] * 4)
    bonus = np.resize(np.array(["0", "1", "2", "3", "4", "Unknown"]), rows)
    rng.shuffle(bonus)
    age = rng.uniform(20 if monitoring else 18, 78 if monitoring else 80, rows)
    exposure = rng.uniform(0.5, 1.5, rows)
    ordered_effect = (
        pd.Series(bonus)
        .map({"0": 0.0, "1": 0.04, "2": 0.12, "3": 0.25, "4": 0.40, "Unknown": 0.18})
        .to_numpy()
    )
    mean_cost = np.exp(5.5 + 0.24 * (region == "North") + ordered_effect + 0.0005 * (age - 45) ** 2)
    mean_cost *= 1.12 if monitoring else 1.0
    dispersion = 35.0
    claim_count = rng.poisson(exposure * np.sqrt(mean_cost) / (0.5 * dispersion))
    total_loss = np.zeros(rows)
    positive = claim_count > 0
    total_loss[positive] = rng.gamma(
        claim_count[positive], 0.5 * dispersion * np.sqrt(mean_cost[positive])
    )
    return pd.DataFrame(
        {
            "policy_id": np.arange(rows) + (10_000 if monitoring else 0),
            "as_of": as_of,
            "region": region,
            "bonus_malus": bonus,
            "driver_age": age,
            "exposure": exposure,
            "burn_cost": total_loss / exposure,
        }
    )


def simulate_demo_deployment(pricing, saved, *, directory: Path, slot: str) -> int:
    """Emulate a published current deployment only in this marked SQLite demo.

    Local publication normally produces LOCAL_AUDIT and cannot deploy through
    the notebook API. This demo setup supplies the SQL state a remote deployment
    would create. It leaves every lineage and immutability guard enabled.
    """
    directory = directory.resolve()
    if pricing.mode != "local" or pricing.engine.dialect.name != "sqlite":
        raise ValueError("the simulated deployment is restricted to local SQLite")
    if not (directory / _MARKER).is_file():
        raise ValueError("the isolated demo directory marker is missing")
    with pricing.engine.begin() as connection:
        databases = connection.exec_driver_sql("PRAGMA database_list").all()
        files = [Path(row[2]).resolve() for row in databases if row[2]]
        if not files or any(not path.is_relative_to(directory) for path in files):
            raise ValueError("all SQLite files must belong to this isolated demo directory")
        existing = connection.execute(
            text("SELECT COUNT(*) FROM pricing.PRICING_MODEL_DEPLOYMENT")
        ).scalar_one()
        if existing:
            raise ValueError("the isolated demo already has a deployment")
        updated = connection.execute(
            text("""UPDATE pricing.PRICING_RATE_PACKAGE SET package_status='PUBLISHED'
                WHERE rate_package_id=:package AND model_id=:model AND package_status='LOCAL_AUDIT'"""),
            {"package": saved.rate_package_id, "model": saved.model_id},
        )
        if updated.rowcount != 1:
            raise ValueError("expected the one local audit package created by this demo")
        result = connection.execute(
            text("""INSERT INTO pricing.PRICING_MODEL_DEPLOYMENT
                (model_id, rate_package_id, deployment_slot, effective_from_ts, deployed_by, deployment_note)
                VALUES (:model, :package, :slot, CURRENT_TIMESTAMP, 'synthetic-demo',
                        'SQLite tutorial simulation; no production deployment')"""),
            {"package": saved.rate_package_id, "model": saved.model_id, "slot": slot},
        )
        return int(result.lastrowid)


_TABLES = (
    ("PRICING_MODEL", "pricing", "model_id", "Registered burn-cost model", "MODEL_RUN.model_id"),
    (
        "MODEL_RECIPE",
        "pricing",
        "recipe_id",
        "Immutable declared feature and estimator configuration",
        "MODEL_RUN.recipe_id",
    ),
    (
        "MODEL_RUN",
        "pricing",
        "model_run_id",
        "Successful build and its artifact/source hashes",
        "model_id; rate_package_id; manifest_id; recipe_id",
    ),
    (
        "PRICING_RATE_PACKAGE",
        "pricing",
        "rate_package_id",
        "Saved rating package",
        "MODEL_RUN.rate_package_id; deployment.rate_package_id",
    ),
    (
        "MODEL_MONITORING_BASELINE",
        "pricing",
        "model_run_id",
        "Immutable SQL snapshot and reference summaries",
        "model_run_id; rate_package_id; source hashes",
    ),
    (
        "PRICING_MODEL_DEPLOYMENT",
        "pricing",
        "deployment_id",
        "Simulated current deployment in the demo slot",
        "model_id; rate_package_id",
    ),
    (
        "DATASET_MANIFEST",
        "pricing",
        "manifest_id",
        "Baseline and later dataset identities and roles",
        "MODEL_RUN.manifest_id; MODEL_MONITOR_RUN.manifest_id",
    ),
    (
        "MODEL_FIT_CONTRACT",
        "mlops",
        "fit_contract_id",
        "Structure and comparison grid shared by monitoring variants",
        "baseline_model_run_id; fit_contract_id",
    ),
    (
        "MODEL_MONITOR_RUN",
        "mlops",
        "variant_code",
        "Four sealed monitoring observations",
        "fit_contract_id; baseline_deployment_id; manifest_id",
    ),
    (
        "MODEL_MONITOR_TERM",
        "mlops",
        "monitor_run_id, sequence_no",
        "Model terms recorded for each monitoring observation",
        "monitor_run_id; term_name",
    ),
    (
        "MODEL_MONITOR_LAMBDA",
        "mlops",
        "monitor_run_id, component_name",
        "Smoothing values and whether they were fixed or estimated",
        "monitor_run_id; term_name; component_name",
    ),
    (
        "MODEL_MONITOR_RELATIVITY",
        "mlops",
        "monitor_run_id, term_name, point_key",
        "Comparable feature multipliers on the baseline grid",
        "monitor_run_id; term_name; point_key",
    ),
    (
        "MODEL_MONITOR_METRIC",
        "mlops",
        "monitor_run_id, metric_name",
        "Fit and score measurements for each observation",
        "monitor_run_id; metric_name",
    ),
)


def export_sql_tables(pricing, *, directory: Path, limit: int = 20) -> tuple[Path, pd.DataFrame]:
    """Export real bounded SQL rows with counts, descriptions, and JSON sidecars."""
    directory = directory.resolve()
    workbook_path = directory / "sql_tables.xlsx"
    if workbook_path.exists():
        raise FileExistsError(workbook_path)
    json_root = directory / "json"
    json_root.mkdir(exist_ok=True)
    summaries, extracts = [], {}
    with pricing.engine.connect() as connection:
        for table, remote_schema, order, description, links in _TABLES:
            count = int(
                connection.execute(text(f"SELECT COUNT(*) FROM pricing.{table}")).scalar_one()
            )
            df = pd.read_sql_query(
                text(f"SELECT * FROM pricing.{table} ORDER BY {order} LIMIT :limit"),
                connection,
                params={"limit": limit},
            )
            for column in df.columns:
                if not column.endswith("_json"):
                    continue
                for index, raw in df[column].items():
                    if not isinstance(raw, str):
                        continue
                    filename = f"{table.lower()}_{index + 1}_{column}.json"
                    (json_root / filename).write_text(
                        json.dumps(json.loads(raw), indent=2, ensure_ascii=False) + "\n"
                    )
                    if len(raw) > 3000:
                        df.at[index, column] = (
                            f"Full JSON: json/{filename}\n\nPreview: {raw[:1500]}\n[preview truncated]"
                        )
            extracts[table] = df
            summaries.append(
                {
                    "SQL Server table": f"{remote_schema}.{table}",
                    "SQLite table / sheet": f"pricing.{table}",
                    "total rows": count,
                    "rows exported": len(df),
                    "extract scope": f"First {limit} rows ordered by {order}"
                    if count > limit
                    else "All rows",
                    "purpose": description,
                    "key links": links,
                }
            )
        stored = connection.execute(
            text(
                "SELECT snapshot_json FROM pricing.MODEL_MONITORING_BASELINE WHERE capture_status='CAPTURED'"
            )
        ).scalar_one()
        snapshot = json.loads(stored)
        (directory / "baseline_snapshot.json").write_text(
            json.dumps(snapshot, indent=2, ensure_ascii=False) + "\n"
        )
    feature_rows = []
    for name in snapshot["recipe"]["feature_order"]:
        feature_rows.append(
            {
                "feature": name,
                "declared configuration": json.dumps(
                    snapshot["recipe"]["features"][name], indent=2
                ),
                "scoring state": json.dumps(snapshot["prediction"]["terms"][name], indent=2),
                "reference profile": json.dumps(snapshot["reference_profiles"].get(name), indent=2),
            }
        )
    guide = pd.DataFrame(summaries)
    readme = pd.DataFrame(
        {
            "item": [
                "Dataset",
                "Deployment",
                "Recurring input",
                "Table locations",
                "Extract scope",
                "JSON",
                "Interpretation",
            ],
            "description": [
                "Synthetic Tweedie burn cost. Baseline 2026-08-31; monitoring 2026-09-15.",
                "One PUBLISHED/current deployment record was simulated only in an isolated SQLite database.",
                "Publication bundle, workbook and receipt files were deleted before loading the monitoring baseline from SQL.",
                "This SQLite mirror places monitoring tables in pricing. Their SQL Server schema is mlops; the baseline table is pricing in both.",
                f"Each table sheet contains at most {limit} actual rows. The tables sheet gives full row counts and sorting.",
                "baseline_snapshot.json and json/ contain complete JSON. Long workbook cells explicitly show a truncated preview.",
                "Monitoring refits are diagnostics; they do not create or deploy new rating packages.",
            ],
        }
    )
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="read_me", index=False)
        guide.to_excel(writer, sheet_name="tables", index=False)
        pd.DataFrame(feature_rows).to_excel(writer, sheet_name="model_features", index=False)
        for table, df in extracts.items():
            df.to_excel(writer, sheet_name=table, index=False)
        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.row_dimensions[1].height = 30
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="16384B")
            for column in worksheet.columns:
                longest = max(len(str(cell.value or "").split("\n")[0]) for cell in column)
                worksheet.column_dimensions[column[0].column_letter].width = min(
                    max(longest + 2, 14), 52
                )
                for cell in column:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            if worksheet.title == "model_features":
                for row in range(2, worksheet.max_row + 1):
                    worksheet.row_dimensions[row].height = 150
    return workbook_path, guide
