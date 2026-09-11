from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from pricing_pipeline.publishing.rating_tables import (
    StagingExport,
    build_staging_frames,
    export_rating_tables,
    model_equivalence_sha256,
    staging_content_sha256,
)


def spline_workbook(path: Path) -> Path:
    w = Workbook()
    s = w.active
    s.title = "Rating Tables"
    s.cell(2, 3, 0.1)
    s.cell(5, 1, "age")
    for col, value in enumerate(["age", "Relativity", "Weight", "a", "b", "c", "d"], 1):
        s.cell(7, col, value)
    rows = [
        ["[-inf, 1e-06)", 1.0, 0.0, 0.0, 0.0, 0.0, 0.0],
        ["[1e-06, 2.0)", 1.0, 20.0, 0.0, 0.2, -0.1, 0.01],
        ["[2.0, inf)", np.exp(0.11), 1.0, 0.11, 0.0, 0.0, 0.0],
    ]
    for row in rows:
        s.append(row)
    w.save(path)
    return path


def frames(path):
    return build_staging_frames(StagingExport(path, "e1", "model", "v1", None, None, {}, "test"))


def test_ppform_import_preserves_coefficients_and_unbounded_tails(tmp_path):
    _, cells, levels = frames(spline_workbook(tmp_path / "rating.xlsx"))
    assert set(cells.term_type) == {"SPLINE_PPOLY_1D"}
    assert cells.spline_b.tolist() == [0.0, 0.2, 0.0]
    assert pd.isna(cells.spline_lower.iloc[0])
    assert cells.spline_lower.iloc[1] == 1e-6
    assert pd.isna(cells.spline_upper.iloc[-1])
    assert set(levels.feature_value_type) == {"NUMERIC"}
    assert set(levels.level_set_type) == {"SPLINE_GRID_1D"}


def test_ppform_coefficient_only_change_changes_both_hashes(tmp_path):
    original = frames(spline_workbook(tmp_path / "rating.xlsx"))
    changed = tuple(df.copy() for df in original)
    changed[1].loc[1, "spline_b"] += 1e-12
    assert staging_content_sha256(*original) != staging_content_sha256(*changed)
    assert model_equivalence_sha256(*original) != model_equivalence_sha256(*changed)


@pytest.mark.parametrize(
    ("cell", "value", "message"),
    [
        ("E7", "missing", "coefficients"),
        ("E9", None, "coefficients"),
        ("E8", 0.1, "tail"),
        ("A9", "[0.0, 2.0)", "contiguous"),
        ("A9", "[2.0, 1.0)", "bounds"),
        ("B9", 2.0, "multiplier"),
    ],
)
def test_ppform_import_rejects_invalid_segments(tmp_path, cell, value, message):
    path = spline_workbook(tmp_path / "rating.xlsx")
    w = load_workbook(path)
    w.active[cell] = value
    w.save(path)
    with pytest.raises(ValueError, match=message):
        frames(path)


def test_ppform_error_domain_preserves_closed_final_endpoint(tmp_path):
    path = spline_workbook(tmp_path / "rating.xlsx")
    w = load_workbook(path)
    w.active.delete_rows(10)
    w.active.delete_rows(8)
    w.active["A8"] = "[1e-06, 2.0]"
    w.save(path)
    _, cells, _ = frames(path)
    assert cells.spline_upper_inclusive.tolist() == [1]


@pytest.mark.parametrize("headers", [[None] * 4, ["q", "r", "s", "t"]])
def test_ppform_rejects_erased_coefficient_headers(tmp_path, headers):
    path = spline_workbook(tmp_path / "rating.xlsx")
    w = load_workbook(path)
    for column, header in enumerate(headers, 4):
        w.active.cell(7, column).value = header
    w.save(path)
    with pytest.raises(ValueError, match="coefficients"):
        frames(path)


def test_ppform_uses_neutral_legacy_cells_for_large_finite_anchors(tmp_path):
    path = spline_workbook(tmp_path / "rating.xlsx")
    w = load_workbook(path)
    for row in range(8, 11):
        w.active.cell(row, 4).value += 21.0
        w.active.cell(row, 2).value = np.exp(w.active.cell(row, 4).value)
    w.save(path)
    _, cells, _ = frames(path)
    assert cells.multiplier.eq(1.0).all()
    assert cells.log_coefficient.eq(0.0).all()
    assert cells.spline_a.iloc[0] == 21.0


@pytest.fixture
def fitted_spline():
    from superglm import Spline, SuperGLM

    x = np.linspace(0, 5, 180)
    df = pd.DataFrame({"age": x})
    y = np.random.default_rng(42).poisson(np.exp(-0.5 + 0.3 * np.sin(x)))
    model = SuperGLM(features={"age": Spline(n_knots=5)}, selection_penalty=0).fit(df, y)
    return model, df, y


@pytest.mark.parametrize("kind", ["ppform", "binned"])
def test_export_preserves_selected_representation(tmp_path, fitted_spline, kind):
    model, df, y = fitted_spline
    path = tmp_path / "rating.xlsx"
    kwargs = {} if kind == "ppform" else {"continuous_kind": kind}
    export_rating_tables(model, df, y, np.ones(len(df)), path, **kwargs)
    exported, cells, _ = frames(path)
    if kind == "binned":
        assert set(cells.term_type) == {"DISCRETIZED_SPLINE_1D"}
        return
    assert set(cells.term_type) == {"SPLINE_PPOLY_1D"}
    points = np.concatenate(
        [df.age, [-1, 6], cells.spline_lower.dropna(), cells.spline_upper.dropna()]
    )
    effects = []
    for x in points:
        matches = cells[
            (cells.spline_lower.isna() | (cells.spline_lower <= x))
            & (
                cells.spline_upper.isna()
                | (x < cells.spline_upper)
                | ((x == cells.spline_upper) & cells.spline_upper_inclusive.eq(1))
            )
        ]
        assert len(matches) == 1
        r = matches.iloc[0]
        u = (
            0.0
            if pd.isna(r.spline_lower) or pd.isna(r.spline_upper)
            else (x - r.spline_lower) / (r.spline_upper - r.spline_lower)
        )
        effects.append(r.spline_a + u * (r.spline_b + u * (r.spline_c + u * r.spline_d)))
    predicted = exported.base_rate.iloc[0] * np.exp(effects)
    np.testing.assert_allclose(predicted, model.predict(pd.DataFrame({"age": points})), rtol=1e-12)
