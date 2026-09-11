from __future__ import annotations

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from superglm import Categorical, SuperGLM

from pricing_pipeline.publishing.rating_tables import export_rating_tables


@pytest.mark.parametrize(
    "levels,kind,expected",
    [
        (["per_unit", "annual"], "discrete", "LOOKUP"),
        (["per_unit"], "discrete", "LOOKUP"),
        (["per_unit", "annual"], "auto", "LOOKUP"),
        (["per_unit"], "auto", "LOOKUP"),
        ([1.0, 2.0], "auto", "LOOKUP"),
        (list(np.linspace(1.0, 3.0, 30)), "auto", "PER_UNIT_FACTOR"),
        ([1.0, 2.0], "per_unit", "PER_UNIT_FACTOR"),
    ],
)
def test_native_offset_export_marks_representation_from_declared_source(
    tmp_path,
    levels,
    kind,
    expected,
):
    source = pd.Series(np.resize(levels, 120), name="Exposure")
    multiplier = (
        source.map({"per_unit": 1.0, "annual": 2.0}).to_numpy(dtype=float)
        if isinstance(levels[0], str)
        else source.to_numpy(dtype=float) / 12.0
    )
    frame = pd.DataFrame({"region": np.resize(["A", "B"], len(source))})
    offset = np.log(multiplier)
    target = np.random.default_rng(913).poisson(2.0 * multiplier)
    model = SuperGLM(
        features={"region": Categorical(base="first")},
        selection_penalty=0.0,
    ).fit(frame, target, offset=offset)
    path = tmp_path / "offset.xlsx"

    export_rating_tables(
        model,
        frame,
        target,
        np.ones(len(frame)),
        path,
        offset=offset,
        offset_source=source,
        offset_kind=kind,
    )

    workbook = load_workbook(path)
    try:
        sheet = workbook["Rating Tables"]
        column = next(
            c for c in range(1, sheet.max_column + 1) if sheet.cell(5, c).value == "Exposure"
        )
        comment = sheet.cell(7, column).comment
        assert comment is not None
        assert comment.text == f"pricing_pipeline:offset_representation={expected}"
        if isinstance(levels[0], str):
            assert [sheet.cell(8 + i, column).value for i in range(len(levels))] == levels
    finally:
        workbook.close()


@pytest.mark.parametrize("level_count,expected", [(2, "LOOKUP"), (30, "PER_UNIT_FACTOR")])
@pytest.mark.parametrize("pass_offset", [True, False])
def test_native_undeclared_offset_export_marks_actual_numeric_representation(
    tmp_path,
    level_count,
    expected,
    pass_offset,
):
    multiplier = np.resize(np.linspace(1.0, 3.0, level_count), 120)
    frame = pd.DataFrame({"region": np.resize(["A", "B"], len(multiplier))})
    offset = np.log(multiplier)
    target = np.random.default_rng(914).poisson(multiplier)
    model = SuperGLM(
        features={"region": Categorical(base="first")},
        selection_penalty=0.0,
    ).fit(frame, target, offset=offset)
    path = tmp_path / "undeclared_offset.xlsx"

    export_rating_tables(
        model,
        frame,
        target,
        np.ones(len(frame)),
        path,
        offset=offset if pass_offset else None,
    )

    workbook = load_workbook(path)
    try:
        sheet = workbook["Rating Tables"]
        column = next(
            c
            for c in range(1, sheet.max_column + 1)
            if sheet.cell(5, c).value == "Offset Multiplier"
        )
        comment = sheet.cell(7, column).comment
        assert comment is not None
        assert comment.text == f"pricing_pipeline:offset_representation={expected}"
    finally:
        workbook.close()
