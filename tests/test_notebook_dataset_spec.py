from dataclasses import replace
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pytest
from sqlalchemy import text
from superglm import Categorical, Numeric, SuperGLM

from pricing_pipeline import notebook as api
from pricing_pipeline.models.config import ModelBuildConfig, ValidationSplitConfig
from pricing_pipeline.resources import migration_root


def _dataset():
    rng = np.random.default_rng(42)
    df = pd.DataFrame(
        {
            "id": np.arange(90),
            "x": rng.uniform(0, 10, 90),
            "exposure": rng.uniform(0.5, 1.5, 90),
            "segment": np.resize(["A", "B", "C"], 90),
            "as_of": ["2026-09-01"] * 90,
        }
    )
    df["y"] = rng.poisson(df["exposure"] * np.exp(0.5 + 0.1 * df["x"]))
    return api.PricingDataset(df, name="test_data", source="test", key="id", as_of="as_of")


def _spec(dataset, **overrides):
    options = {
        "name": "TEST_MODEL",
        "label": "Test model",
        "target": "y",
        "model_type": "superglm_poisson",
        "deployment_slot": "TEST_UAT",
        "dataset": dataset,
        "features": ["log_x", "segment"],
        "transforms": {"log_x": api.Log1p("x"), "log_exposure": api.Log("exposure")},
        "offset_column": "log_exposure",
        "export_weight_column": "exposure",
        "fit_mode": "fit",
        "validation": ValidationSplitConfig.kfold(n_splits=2, random_state=42),
    }
    options.update(overrides)
    return api.PricingModelSpec(**options)


def test_dataset_spec_derives_data_and_offset_fields():
    assert hasattr(api, "PricingDataset"), "the approved dataset helper must be public"
    dataset = _dataset()
    spec = _spec(dataset)
    assert (spec.dataset_name, spec.source_system, spec.pk_columns, spec.data_as_of_column) == (
        "test_data",
        "test",
        ("id",),
        "as_of",
    )
    assert spec.features == ("log_x", "segment")
    assert spec.offset_source_column == "exposure"
    assert spec.offset_label == "log(exposure)"
    assert replace(spec, scoring=("deviance",)).dataset is dataset


def test_spec_defaults_to_exact_splines_and_validates_band_opt_in():
    spec = _spec(_dataset())
    assert spec.spline_export == "exact"
    assert replace(spec, spline_export="binned").spline_export == "binned"
    with pytest.raises(ValueError, match="spline_export"):
        replace(spec, spline_export="unknown")


@pytest.mark.parametrize(
    "override",
    [
        {"dataset_name": "other"},
        {"source_system": "other"},
        {"pk_columns": ("x",)},
        {"data_as_of_column": "other"},
        {"offset_source_column": "x"},
        {"offset_label": "log1p(exposure)"},
    ],
)
def test_dataset_spec_rejects_conflicting_metadata(override):
    with pytest.raises(ValueError, match="conflict|match"):
        _spec(_dataset(), **override)


@pytest.mark.parametrize("features", [{"log_x", "segment"}, "log_x"])
def test_dataset_spec_rejects_unordered_or_scalar_feature_names(features):
    with pytest.raises((TypeError, ValueError), match="ordered|sequence"):
        _spec(_dataset(), features=features)


@pytest.mark.parametrize("change", ["target", "source", "transform", "row_order"])
def test_build_checks_dataset_and_transforms_before_using_database(tmp_path, change):
    dataset = _dataset()
    spec = _spec(dataset)
    df = api.apply_transforms(dataset.df, spec.transforms)
    if change == "target":
        df.loc[0, "y"] += 1
    elif change == "source":
        df.loc[0, "x"] += 1
    elif change == "transform":
        df.loc[0, "log_x"] += 1
    else:
        df = df.iloc[::-1]
    model = api.RegisteredModel(
        model_id=1,
        config=ModelBuildConfig(
            model_name=spec.name,
            model_label=spec.label,
            target_name=spec.target,
            model_type=spec.model_type,
            deployment_slot=spec.deployment_slot,
            validation_split=spec.validation,
        ),
        source_root=tmp_path,
        spec=spec,
    )
    context = api.NotebookContext(
        engine=object(),
        settings=SimpleNamespace(),
        mode="local",
        write_allowed=True,
        destination="unused",
    )
    with pytest.raises(ValueError, match="prepared|dataset|transform"):
        api.fit_model(context, model=model, frame=df, superglm_model=object())


@pytest.mark.parametrize("retain_fit_state", [True, False])
def test_transformed_dataset_build_publishes_preparation_metadata(tmp_path, retain_fit_state):
    import json

    from openpyxl import load_workbook

    dataset = _dataset()
    dataset_path = tmp_path / "dataset.joblib"
    dataset.save(dataset_path)
    spec = _spec(api.PricingDataset.load(dataset_path))
    df = api.apply_transforms(spec.dataset.df, spec.transforms)
    source_root = tmp_path / "model"
    source_root.mkdir()
    (source_root / "model.py").write_text("# Test model source\n")
    pricing = api.connect(mode="local", local_root=tmp_path / "local")
    try:
        model = api.register_model(pricing, spec, source_root=source_root)
        candidate = api.fit_model(
            pricing,
            model=model,
            frame=df,
            superglm_model=SuperGLM(
                family="poisson",
                selection_penalty=0.0,
                retain_fit_state=retain_fit_state,
                features={"log_x": Numeric(), "segment": Categorical()},
            ),
        )
        published = api.save_model_version(pricing, candidate)
        with pricing.engine.connect() as connection:
            metadata = json.loads(
                connection.execute(
                    text(
                        "SELECT package_metadata_json FROM pricing.PRICING_RATE_PACKAGE "
                        "WHERE rate_package_id = :id"
                    ),
                    {"id": published.rate_package_id},
                ).scalar_one()
            )
        prep = metadata["input_preparation"]
        assert prep["scoring_input"] == "prepared"
        assert prep["transforms"]["log_x"] == {"source": "x", "operation": "log1p"}
        assert prep["transforms"]["log_exposure"] == {"source": "exposure", "operation": "log"}
        workbook = load_workbook(candidate.completed_build.rating_workbook_path, read_only=True)
        try:
            assert "Input Preparation" in workbook.sheetnames
            values = repr(list(workbook["Input Preparation"].values))
            assert "log_x" in values and "log_exposure" in values
        finally:
            workbook.close()
    finally:
        pricing.engine.dispose()


@pytest.mark.parametrize("spline_export", ["exact", "binned"])
@pytest.mark.parametrize("retain_fit_state", [True, False])
@pytest.mark.parametrize("fit_mode", ["fit", "fit_reml"])
def test_transformed_spline_build_keeps_export_choice_and_sql_curve(
    tmp_path, spline_export, retain_fit_state, fit_mode
):
    import joblib
    from superglm import Spline

    dataset = _dataset()
    spec = _spec(dataset, spline_export=spline_export, fit_mode=fit_mode)
    df = api.apply_transforms(dataset.df, spec.transforms)
    source_root = tmp_path / "model"
    source_root.mkdir()
    (source_root / "model.py").write_text("# Spline publication integration test\n")
    pricing = api.connect(mode="local", local_root=tmp_path / "local")
    try:
        model = api.register_model(pricing, spec, source_root=source_root)
        candidate = api.fit_model(
            pricing,
            model=model,
            frame=df,
            superglm_model=SuperGLM(
                family="poisson",
                selection_penalty=0.0,
                retain_fit_state=retain_fit_state,
                features={"log_x": Spline(n_knots=5), "segment": Categorical()},
            ),
        )
        bundle = joblib.load(candidate.completed_build.candidate_artifact_path)["bundle"]
        if not retain_fit_state:
            assert bundle.fitted_model._dm is None
            assert bundle.fitted_model._fit_X_ref is None
            assert bundle.fitted_model.summary() is not None
        assert bundle.continuous_kind == ("ppform" if spline_export == "exact" else "binned")
        published = api.save_model_version(pricing, candidate)
        with pricing.engine.connect() as connection:
            metrics = {
                row.metric_name: (row.metric_value, row.metric_scope)
                for row in connection.execute(
                    text(
                        "SELECT metric_name, metric_value, metric_scope "
                        "FROM mlops.MODEL_RUN_METRIC WHERE model_run_id = "
                        "(SELECT model_run_id FROM pricing.MODEL_RUN WHERE rate_package_id=:id)"
                    ),
                    {"id": published.rate_package_id},
                )
            }
            telemetry = bundle.cv_report["full_fit_telemetry"]
            assert metrics["fit_deviance"] == (
                pytest.approx(telemetry["fit"]["deviance"]),
                "full_fit",
            )
            assert metrics["fit_effective_df"] == (
                pytest.approx(telemetry["fit"]["effective_df"]),
                "full_fit",
            )
            assert metrics["fit_converged"] == (1.0, "full_fit")
            assert metrics["fit_n_obs"] == (len(df), "full_fit")
            assert metrics["fit_reml_enabled"] == (float(fit_mode == "fit_reml"), "full_fit")
            assert metrics["cv_pooled_deviance"][1] == "cv"
            if fit_mode == "fit_reml":
                assert metrics["fit_reml_converged"] == (
                    float(telemetry["reml"]["converged"]),
                    "full_fit",
                )
                assert metrics["fit_reml_n_iter"][0] == telemetry["reml"]["n_reml_iter"]
            else:
                assert "fit_reml_converged" not in metrics
            # Run the production view query against the attached local databases.
            # COUNT_BIG is the only dialect change; this does not exercise SQL Server DDL.
            migration = migration_root() / "V046__full_fit_diagnostics.sql"
            assert migration.exists(), "SQL Server must expose the saved fit diagnostics"
            query = "WITH run_metrics" + migration.read_text().split("WITH run_metrics", 1)[1]
            query = query.split(";", 1)[0].replace("COUNT_BIG(", "COUNT(")
            summary = connection.execute(text(query)).mappings().one()
            assert summary["fit_deviance"] == pytest.approx(telemetry["fit"]["deviance"])
            assert summary["fit_effective_df"] == pytest.approx(telemetry["fit"]["effective_df"])
            assert summary["pooled_deviance"] == pytest.approx(metrics["cv_pooled_deviance"][0])
            assert summary["fit_reml_enabled"] == float(fit_mode == "fit_reml")
            # Old runs and incorrectly scoped metrics must not acquire fit evidence.
            connection.execute(
                text(
                    "UPDATE mlops.MODEL_RUN_METRIC SET metric_scope='cv' "
                    "WHERE metric_name='fit_deviance'"
                )
            )
            assert connection.execute(text(query)).mappings().one()["fit_deviance"] is None
            connection.execute(
                text("DELETE FROM mlops.MODEL_RUN_METRIC WHERE metric_scope='full_fit'")
            )
            old_summary = connection.execute(text(query)).mappings().one()
            assert old_summary["fit_effective_df"] is None
            assert old_summary["pooled_deviance"] is not None
            segments = pd.read_sql_query(
                text(
                    "SELECT * FROM pricing.V_MODEL_SPLINE_SEGMENT WHERE rate_package_id=:id "
                    "ORDER BY segment_order"
                ),
                connection,
                params={"id": published.rate_package_id},
            )
            if spline_export == "binned":
                assert segments.empty
                return
            assert not segments.empty
            assert segments.transforms_json.str.contains("log1p").all()
            assert segments.term_type.eq("SPLINE_PPOLY_1D").all()
            assert (
                connection.execute(
                    text(
                        "SELECT COUNT(*) FROM pricing.V_MODEL_RELATIVITY "
                        "WHERE rate_package_id=:id AND term_name='log_x'"
                    ),
                    {"id": published.rate_package_id},
                ).scalar_one()
                == 0
            )
            # Evaluate persisted coefficients with a SQL expression at a knot,
            # both tails, and an interior input. No rounding of the spline grid.
            for x in [
                segments.lower_bound.dropna().min() - 1,
                segments.lower_bound.dropna().iloc[1],
                float(df.log_x.median()),
                segments.upper_bound.dropna().max() + 1,
            ]:
                effect = connection.execute(
                    text("""
                    SELECT a + u*(b + u*(c + u*d)) FROM (
                        SELECT a,b,c,d,
                            CASE WHEN lower_bound IS NULL OR upper_bound IS NULL THEN 0.0
                            ELSE (:x-lower_bound)/(upper_bound-lower_bound) END AS u
                        FROM pricing.PRICING_SPLINE_SEGMENT
                        WHERE rate_package_id=:id AND feature_name='log_x'
                        AND (lower_bound IS NULL OR :x>=lower_bound)
                        AND (upper_bound IS NULL OR :x<upper_bound
                             OR (:x=upper_bound AND upper_inclusive=1))
                    )
                """),
                    {"x": float(x), "id": published.rate_package_id},
                ).scalar_one()
                # Hold the category and offset fixed to isolate the spline.
                probe = pd.DataFrame(
                    {"log_x": [x, float(df.log_x.median())], "segment": ["A", "A"]}
                )
                pred = bundle.fitted_model.predict(probe, offset=np.zeros(2))
                # Reference effect is evaluated from stored segments as well.
                r = segments[
                    (segments.lower_bound.isna() | (segments.lower_bound <= probe.log_x[1]))
                    & (segments.upper_bound.isna() | (probe.log_x[1] < segments.upper_bound))
                ].iloc[0]
                u = (probe.log_x[1] - r.lower_bound) / (r.upper_bound - r.lower_bound)
                mid_effect = r.a + u * (r.b + u * (r.c + u * r.d))
                assert np.exp(effect - mid_effect) == pytest.approx(pred[0] / pred[1], rel=1e-12)
    finally:
        pricing.engine.dispose()
