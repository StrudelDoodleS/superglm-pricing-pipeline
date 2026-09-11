from __future__ import annotations

import hashlib
from dataclasses import replace
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from superglm import Numeric, SuperGLM

from pricing_pipeline.modeling.monitoring import build_model_fit_contract, run_monitoring_fit
from pricing_pipeline.publishing import editor
from pricing_pipeline.publishing.metadata import (
    OffsetExportContract,
    build_superglm_publication_receipt,
    canonical_receipt_bytes,
    load_publication_receipt,
    write_publication_receipt,
)
from pricing_pipeline.publishing.rating_tables import (
    export_rating_tables,
    rating_workbook_model_equivalence_sha256,
)
from pricing_pipeline.workbench.artifacts import (
    CandidateBundle,
    load_candidate_bundle,
    save_candidate_bundle,
)


@pytest.fixture
def fitted():
    X = pd.DataFrame({"log_density": np.linspace(0.0, 5.0, 80)})
    y = np.random.default_rng(31).poisson(np.exp(-0.8 + 0.2 * X["log_density"]))
    model = SuperGLM(features={"log_density": Numeric()}, selection_penalty=0.0).fit(X, y)
    return model, X, y


def _transforms():
    return {"log_density": {"operation": "log", "source": "density"}}


def _receipt(model, transforms=None):
    return build_superglm_publication_receipt(
        model, offset_contract=OffsetExportContract(handling="NONE"), input_transforms=transforms
    )


def test_receipt_persists_input_preparation_and_changes_equivalence(fitted, tmp_path):
    model, X, y = fitted
    metadata = _transforms()
    receipt = _receipt(model, metadata)
    metadata["log_density"]["source"] = "changed_after_receipt"
    assert receipt.model_dump(mode="json")["package_metadata"]["input_preparation"] == {
        "schema_version": 1,
        "scoring_input": "prepared",
        "transforms": _transforms(),
    }
    legacy = _receipt(model)
    assert "input_preparation" not in legacy.package_metadata
    assert canonical_receipt_bytes(legacy) == canonical_receipt_bytes(_receipt(model, {}))
    workbook = export_rating_tables(model, X, y, None, tmp_path / "rating.xlsx")
    digests = []
    for index, current in enumerate((legacy, receipt)):
        path = tmp_path / f"receipt-{index}.json"
        digest = write_publication_receipt(current, path)
        assert load_publication_receipt(path, digest) == current
        digests.append(
            rating_workbook_model_equivalence_sha256(
                workbook_path=workbook,
                export_id="export-1",
                model_name="TEST",
                model_version="v1",
                effective_from="2026-09-11",
                publication_receipt_path=path,
                publication_receipt_sha256=digest,
            )
        )
    assert digests[0] != digests[1]


def test_workbook_documents_preparation_without_changing_rating_units(fitted, tmp_path):
    model, X, y = fitted
    plain = export_rating_tables(model, X, y, None, tmp_path / "plain.xlsx")
    prepared = export_rating_tables(
        model, X, y, None, tmp_path / "prepared.xlsx", input_transforms=_transforms()
    )
    with_plain = load_workbook(plain)
    with_preparation = load_workbook(prepared)
    try:
        for sheet in with_plain:
            assert list(sheet.values) == list(with_preparation[sheet.title].values)
        rows = list(with_preparation["Input Preparation"].values)
        cells = [str(value) for row in rows for value in row if value is not None]
        assert "log_density" in cells
        assert "density" in cells
        assert "log" in cells
        assert "log(density)" in cells
        assert (
            "Prepare model feature values before SQL scoring. Rating-table units are unchanged."
            in cells
        )
    finally:
        with_plain.close()
        with_preparation.close()


def _bundle(fitted, **kwargs):
    model, X, y = fitted
    return CandidateBundle(
        fitted_model=model,
        X=X,
        y=y,
        sample_weight=None,
        offset=None,
        export_weight=None,
        cv_report={},
        model_name="TEST",
        model_version="v1",
        export_id="export-1",
        manifest_id="manifest-1",
        split_set_id=None,
        pk_columns=("id",),
        row_order_sha256="a" * 64,
        model_source_sha256="b" * 64,
        model_frame_sha256="c" * 64,
        offset_contract=OffsetExportContract(handling="NONE"),
        **kwargs,
    )


def _roundtrip(bundle, tmp_path):
    path = tmp_path / "candidate.joblib"
    metadata = save_candidate_bundle(bundle, path)
    return load_candidate_bundle(
        path,
        expected_sha256=metadata.sha256,
        expected_size_bytes=metadata.size_bytes,
        expected_format=metadata.format,
        expected_python_version=metadata.python_version,
        expected_superglm_version=metadata.superglm_version,
        allowed_root=tmp_path,
    )


def test_bundle_copies_metadata_and_preserves_it_on_reload_and_edit(fitted, tmp_path):
    metadata = _transforms()
    bundle = _bundle(fitted, input_transforms=metadata)
    metadata["log_density"]["source"] = "mutated"
    loaded = _roundtrip(replace(bundle, export_id="edited-1"), tmp_path)
    assert loaded.input_transforms == _transforms()
    assert loaded.export_id == "edited-1"


def test_legacy_bundle_without_transform_field_loads(fitted, tmp_path):
    bundle = _bundle(fitted)
    bundle.__dict__.pop("input_transforms", None)
    assert _roundtrip(bundle, tmp_path).input_transforms is None


def test_monitoring_contract_preserves_prepared_input_semantics(fitted):
    model, X, y = fitted
    contract = build_model_fit_contract(model, input_transforms=_transforms())
    result = run_monitoring_fit(model, X, y, variant="STATIC_SCORE", input_transforms=_transforms())
    assert result.contract.contract_sha256 == contract.contract_sha256
    assert contract.payload()["structure"]["package_metadata"]["input_preparation"] == {
        "schema_version": 1,
        "scoring_input": "prepared",
        "transforms": _transforms(),
    }
    assert contract.contract_sha256 != build_model_fit_contract(model).contract_sha256


@pytest.mark.parametrize("invalid", [{"x": {"operation": "eval", "source": "raw"}}, {"x": str}])
def test_publication_rejects_unsupported_transform_metadata(fitted, tmp_path, invalid):
    model, X, y = fitted
    with pytest.raises((TypeError, ValueError), match="transform|operation"):
        _receipt(model, invalid)
    path = tmp_path / "invalid.xlsx"
    with pytest.raises((TypeError, ValueError), match="transform|operation"):
        export_rating_tables(model, X, y, None, path, input_transforms=invalid)
    assert not path.exists()


@pytest.mark.parametrize("model_kind", ["EDITOR_EDIT", "MANUAL_EDIT"])
def test_edited_export_inherits_parent_preparation(fitted, tmp_path, model_kind):
    bundle = _bundle(fitted, input_transforms=_transforms())
    parent = SimpleNamespace(
        model_id=1,
        model_name="TEST",
        model_version="v1",
        effective_from=None,
        config=SimpleNamespace(model_type="superglm_poisson", target_name="claims"),
        bundle=bundle,
        champion=editor.ChampionSnapshot(
            deployment_slot="TEST_UAT",
            rate_package_id=None,
            bundle=None,
            unavailable_reason="No champion",
        ),
    )
    submission = SimpleNamespace(
        model_kind=model_kind,
        submission_id="edit-1",
        deployment_slot="TEST_UAT",
        manifest_id=bundle.manifest_id,
        split_set_id=None,
        model_source_sha256=bundle.model_source_sha256,
        reason="Test preparation inheritance",
        claimed_identity="analyst",
        parent_rate_package_id=1,
        parent_model_run_id=1,
        path=str(tmp_path / "submission.json"),
        sha256="c" * 64,
        editor_session_path=str(tmp_path / "session.json"),
        editor_session_sha256="d" * 64,
        editor_session_size_bytes=10,
        baseline_candidate_sha256="e" * 64,
    )
    result = editor.export_edited_model(
        parent,
        submission,
        created_by="publisher",
        allowed_root=tmp_path,
        write_dir=tmp_path,
        published_dir=tmp_path,
        edited_model=fitted[0],
    )
    receipt = load_publication_receipt(
        result.completed_build.publication_receipt_path,
        result.completed_build.publication_receipt_sha256,
    )
    assert receipt.package_metadata["input_preparation"]["transforms"] == _transforms()
    assert result.bundle.input_transforms == _transforms()
    workbook = load_workbook(result.completed_build.rating_workbook_path)
    try:
        assert workbook["Input Preparation"]["A5"].value == "log_density"
    finally:
        workbook.close()


def test_monitoring_rebuilds_receipt_from_verified_bundle_preparation(fitted, tmp_path):
    from pricing_pipeline.modeling.monitoring import MonitoringError
    from tests.test_model_monitoring import _monitoring_candidate

    model, X, y = fitted
    candidate = _monitoring_candidate(tmp_path, model, X, y)
    bundle = replace(candidate.bundle, input_transforms=_transforms())
    artifact = save_candidate_bundle(bundle, tmp_path / "prepared-candidate.joblib")
    receipt_digest = hashlib.sha256(
        canonical_receipt_bytes(_receipt(model, _transforms()))
    ).hexdigest()
    technical = {
        **candidate.technical,
        "candidate_artifact_path": artifact.path,
        "candidate_artifact_sha256": artifact.sha256,
        "candidate_artifact_size_bytes": artifact.size_bytes,
        "publication_receipt_sha256": receipt_digest,
        "package_publication_receipt_sha256": receipt_digest,
    }
    candidate = replace(candidate, bundle=bundle, technical=technical)
    candidate.workbench.open = lambda *args, **kwargs: candidate
    # The verified artifact is authoritative even when the in-memory bundle was modified.
    bundle.input_transforms["log_density"]["source"] = "untrusted_memory"
    result = run_monitoring_fit(candidate, X, y, variant="STATIC_SCORE")
    preparation = result.contract.payload()["structure"]["package_metadata"]["input_preparation"]
    assert preparation["transforms"] == _transforms()
    with pytest.raises(MonitoringError, match="input_transforms does not match"):
        run_monitoring_fit(candidate, X, y, variant="STATIC_SCORE", input_transforms={})


def test_preparation_sheet_keeps_formula_looking_column_names_as_text(fitted, tmp_path):
    model, X, y = fitted
    path = export_rating_tables(
        model,
        X,
        y,
        None,
        tmp_path / "formula-name.xlsx",
        input_transforms={"=model_column": {"operation": "log", "source": "=source_column"}},
    )
    workbook = load_workbook(path)
    try:
        sheet = workbook["Input Preparation"]
        assert sheet["A5"].value == "=model_column"
        assert sheet["B5"].value == "=source_column"
        assert sheet["A5"].data_type == "s"
        assert sheet["B5"].data_type == "s"
    finally:
        workbook.close()
