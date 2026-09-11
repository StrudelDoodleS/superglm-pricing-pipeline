"""Compare active procedure queries with native fitted-model predictions.

SQLite executes the production INSERT/SELECT expressions after explicit dialect
translations for TRY_CONVERT, scalar CROSS APPLY, and CONVERT. It does not execute
T-SQL control flow, SQL Server DECIMAL storage, or its query planner. The fixture
uses exact splines, numeric terms, categories and a continuous exported offset;
the separate legacy band tests cover band selection.
"""

import ast
import hashlib
import inspect
import json
import math
import re
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.comments import Comment
from sqlalchemy import text
from superglm import Categorical, Numeric, Spline, SuperGLM

from pricing_pipeline.data.transforms import Log1p, apply_transforms, transforms_from_metadata
from pricing_pipeline.infra.offline_sqlite import (
    apply_offline_ddl,
    sqlite_engine_with_offline_schemas,
)
from pricing_pipeline.publishing.metadata import (
    OffsetExportContract,
    build_superglm_publication_receipt,
    write_publication_receipt,
)
from pricing_pipeline.publishing.rating_tables import _verified_rating_frames, export_rating_tables
from pricing_pipeline.publishing.sqlserver import _insert_rating_tables
from pricing_pipeline.resources import migration_root


def _active_procedure():
    result = None
    for migration in sorted(migration_root().glob("V*.sql")):
        match = re.search(
            r"CREATE OR ALTER PROCEDURE pricing\.PREDICT_RATE_PACKAGE\b.*?(?=\nGO|\Z)",
            migration.read_text(),
            re.DOTALL,
        )
        if match:
            result = match[0]
    assert result is not None
    return result


def _float_or_none(value):
    try:
        number = float(value)
        return number if math.isfinite(number) else None
    except ValueError, TypeError:
        return None


def _json_value(document, path):
    if document is None:
        return None
    value = json.loads(document).get(path[2:])
    if value is None or isinstance(value, (dict, list)):
        return None
    if isinstance(value, bool):
        return str(value).lower()
    return str(value)


def _sqlite_query(query):
    """Only change dialect syntax; retain production predicates and arithmetic."""
    apply = re.search(
        r"CROSS APPLY \(\s*SELECT (.*?)\s*\) AS raw_input\s*"
        r"CROSS APPLY \(\s*SELECT (.*?)\s*\) AS numeric_input",
        query,
        re.DOTALL,
    )
    if apply:
        raw = (
            "SELECT term_feature.term_id AS join_term_id, "
            + apply[1]
            + " FROM pricing.PRICING_TERM_FEATURE AS term_feature WHERE position_no = 1"
        )
        numeric = "SELECT raw_input.join_term_id, " + apply[2] + " FROM (" + raw + ") AS raw_input"
        query = (
            query[: apply.start()]
            + (
                "JOIN (" + raw + ") AS raw_input ON raw_input.join_term_id = cell.term_id "
                "JOIN ("
                + numeric
                + ") AS numeric_input ON numeric_input.join_term_id = cell.term_id "
            )
            + query[apply.end() :]
        )
    query = re.sub(r"TRY_CONVERT\(FLOAT,\s*([\w.]+)\)", r"TRY_FLOAT(\1)", query)
    query = query.replace("CONVERT(NVARCHAR(128), segment_order)", "CAST(segment_order AS TEXT)")
    query = re.sub(r"@matched\b", "matched", query)
    assert "CROSS APPLY" not in query
    assert "TRY_CONVERT" not in query
    return query


def _run_procedure_queries(connection, features, *, package=1, exposure=1):
    raw = connection.connection.driver_connection
    sql = _active_procedure()
    raw.execute("DROP TABLE IF EXISTS matched")
    declaration = re.search(r"DECLARE @matched TABLE (\(.*?\));", sql, re.DOTALL)[1]
    raw.execute("CREATE TEMP TABLE matched " + declaration)
    params = {
        "features_json": json.dumps(features),
        "rate_package_id": package,
        "exposure": exposure,
    }
    queries = re.findall(r"WITH spline_input AS \(.*?;|INSERT INTO @matched \(.*?;", sql, re.DOTALL)
    assert queries
    for query in queries:
        if "'BAND'" in query:
            # Native fixture has no binned effects; skipping its SQL Server TOP/APPLY is explicit.
            assert (
                raw.execute(
                    "SELECT COUNT(*) FROM pricing.PRICING_COMPILED_1D_RATE_BAND"
                ).fetchone()[0]
                == 0
            )
            continue
        raw.execute(_sqlite_query(query), params)
    count_sql = re.search(r"SELECT @required_terms = (.*?);", sql, re.DOTALL)[1]
    required = raw.execute("SELECT " + count_sql, params).fetchone()[0]
    matched = raw.execute("SELECT COUNT(*) FROM matched").fetchone()[0]
    # Mirror only the procedure's THROW branch, which SQLite cannot execute.
    assert "IF @matched_terms <> @required_terms" in sql
    if matched != required:
        names = raw.execute("SELECT term_name, match_type FROM matched ORDER BY term_id").fetchall()
        raise ValueError(f"Input features did not match every required term: {names}")
    package_row = raw.execute(
        "SELECT model_name, base_rate FROM pricing.PRICING_RATE_PACKAGE WHERE rate_package_id = ?",
        (package,),
    ).fetchone()
    params.update(
        model_name=package_row[0],
        base_rate=package_row[1],
        required_terms=required,
        matched_terms=matched,
    )
    summary = re.search(r"SELECT\s+@model_name AS model_name,.*?FROM @matched;", sql, re.DOTALL)[0]
    result = raw.execute(_sqlite_query(summary), params).fetchone()
    return result[5], raw.execute(
        "SELECT term_name, match_type, log_coefficient FROM matched ORDER BY term_id"
    ).fetchall()


def _stage(connection, table, frame):
    frame = frame.astype(object).where(pd.notna(frame), None)
    columns = list(frame.columns)
    query = (
        f"INSERT INTO pricing_stg.{table} ("
        + ", ".join(columns)
        + ") VALUES ("
        + ", ".join(":" + col for col in columns)
        + ")"
    )
    connection.execute(text(query), frame.to_dict("records"))


def _insert_frames_package(connection, frames, package):
    connection.execute(
        text("""INSERT INTO pricing.PRICING_RATE_PACKAGE
        (rate_package_id, model_id, model_name, model_version, package_version, base_rate, package_metadata_json, package_status, created_by)
        VALUES (:package, 1, 'NATIVE', :version, :package, :base, :metadata, 'DRAFT', 'pytest')"""),
        {
            "package": package,
            "version": f"v{package}",
            "base": float(frames[0].base_rate.iloc[0]),
            "metadata": frames[0].package_metadata_json.iloc[0],
        },
    )
    for table, frame in zip(
        ("STG_RATING_EXPORT", "STG_RATE_CELL", "STG_CELL_LEVEL", "STG_TERM_METADATA"),
        frames,
        strict=True,
    ):
        _stage(connection, table, frame)
    # Execute the publisher's actual relational INSERT SQL, splitting only SQLAlchemy's batch.
    tree = ast.parse(inspect.getsource(_insert_rating_tables))
    insert_sql = next(
        node.value
        for node in ast.walk(tree)
        if isinstance(node, ast.Constant)
        and isinstance(node.value, str)
        and "INSERT INTO pricing.PRICING_SPLINE_SEGMENT" in node.value
    )
    for statement in insert_sql.split(";"):
        if statement.strip():
            connection.execute(
                text(statement),
                {
                    "model_id": 1,
                    "rate_package_id": package,
                    "export_id": frames[0].export_id.iloc[0],
                },
            )


@pytest.fixture(params=["clip", "error"])
def native_package(request, tmp_path):
    rng = np.random.default_rng(67)
    raw_frame = pd.DataFrame(
        {
            "raw_density": np.expm1(np.linspace(0.2, 4.5, 240)),
            "numeric": rng.uniform(-1, 2, 240),
            "region": np.resize(["A", "B", "C"], 240),
        }
    )
    X = apply_transforms(raw_frame, {"density": Log1p("raw_density")})
    exposure = pd.Series(rng.uniform(0.1, 3, len(X)), name="Exposure")
    scale = 0.2  # Nonunit conversion catches confusing a per-unit factor with a power term.
    offset = np.log(exposure * scale)
    y = rng.poisson(
        np.exp(0.4 + 0.3 * np.sin(X.density) + 0.25 * X.numeric + 0.2 * X.region.eq("B") + offset)
    )
    model = SuperGLM(
        features={
            "density": Spline(n_knots=5, extrapolation=request.param),
            "numeric": Numeric(),
            "region": Categorical(),
        },
        selection_penalty=0,
    ).fit(X, y, offset=offset)
    workbook = export_rating_tables(
        model,
        X,
        y,
        np.ones(len(X)),
        tmp_path / "native.xlsx",
        offset=offset,
        offset_source=exposure,
        offset_name="Exposure",
        offset_kind="per_unit",
        input_transforms={"density": {"operation": "log1p", "source": "raw_density"}},
    )
    receipt = build_superglm_publication_receipt(
        model,
        offset_contract=OffsetExportContract(
            handling="EXPORTED_FACTOR",
            source_factor_name="Exposure",
            published_factor_name="Exposure",
            source_name="Exposure",
            label="log(Exposure * 0.2)",
        ),
        input_transforms={"density": {"operation": "log1p", "source": "raw_density"}},
    )
    receipt_path = tmp_path / "receipt.json"
    digest = write_publication_receipt(receipt, receipt_path)
    frames = _verified_rating_frames(
        workbook_path=workbook,
        export_id="native",
        model_name="NATIVE",
        model_version="v1",
        effective_from=None,
        publication_receipt_path=receipt_path,
        publication_receipt_sha256=digest,
    )
    engine = sqlite_engine_with_offline_schemas(
        {name: tmp_path / f"{name}.sqlite" for name in ("pricing", "pricing_stg", "mlops")}
    )
    apply_offline_ddl(engine)
    with engine.begin() as connection:
        raw = connection.connection.driver_connection
        raw.create_function(
            "HASHBYTES", 2, lambda algorithm, value: hashlib.sha256(value.encode()).digest()
        )
        raw.create_function("TRY_FLOAT", 1, _float_or_none)
        raw.create_function("JSON_VALUE", 2, _json_value)
        raw.create_function(
            "CONCAT", -1, lambda *args: "".join("" if arg is None else str(arg) for arg in args)
        )
        raw.create_function("LOG", 1, lambda value: None if value is None else math.log(value))
        raw.create_function("EXP", 1, lambda value: None if value is None else math.exp(value))
        connection.execute(
            text("""INSERT INTO pricing.PRICING_MODEL
            (model_id, model_name, model_label, target_name, model_type, created_by)
            VALUES (1, 'NATIVE', 'Native', 'y', 'superglm_poisson', 'pytest')""")
        )
        _insert_frames_package(connection, frames, 1)
        yield SimpleNamespace(
            connection=connection,
            model=model,
            X=X,
            y=y,
            receipt_path=receipt_path,
            receipt_sha256=digest,
            raw_frame=raw_frame,
            exposure=exposure,
            scale=scale,
            boundary=request.param,
        )
    engine.dispose()


def test_native_mixed_model_predictions_match_active_sql(native_package):
    fixture = native_package
    term_metadata = fixture.connection.execute(
        text(
            "SELECT term_metadata_json FROM pricing.PRICING_TERM WHERE term_type = 'OFFSET_FACTOR'"
        )
    ).scalar_one()
    assert json.loads(term_metadata)["rating_representation"] == "PER_UNIT_FACTOR"
    knots = (
        fixture.connection.execute(
            text("""SELECT lower_bound FROM pricing.PRICING_SPLINE_SEGMENT WHERE lower_bound IS NOT NULL
        UNION SELECT upper_bound FROM pricing.PRICING_SPLINE_SEGMENT WHERE upper_bound IS NOT NULL ORDER BY 1""")
        )
        .scalars()
        .all()
    )
    adjacent = [
        point
        for knot in knots[1:-1]
        for point in (np.nextafter(knot, -np.inf), np.nextafter(knot, np.inf))
    ]
    extra = (
        knots + adjacent + ([knots[0] - 0.1, knots[-1] + 0.5] if fixture.boundary == "clip" else [])
    )
    X = pd.concat(
        [fixture.X.iloc[::17], pd.DataFrame({"density": extra, "numeric": 0.7, "region": "B"})],
        ignore_index=True,
    )
    # Include fractional factors and values above the training range.
    exposure = np.resize([0.5, 0.001, 1.0, 3.7, 1000.0], len(X))
    expected = fixture.model.predict(X, offset=np.log(exposure * fixture.scale))
    metadata = fixture.connection.execute(
        text("SELECT transforms_json FROM pricing.V_MODEL_SPLINE_SEGMENT LIMIT 1")
    ).scalar_one()
    raw_inputs = X.drop(columns="density").copy()
    raw_inputs["raw_density"] = np.expm1(X.density)
    prepared = apply_transforms(raw_inputs, transforms_from_metadata(json.loads(metadata)))
    # Keep exact exported knot bit patterns after separately checking transform reconstruction.
    np.testing.assert_allclose(prepared.density, X.density, rtol=1e-15, atol=1e-15)
    actual = []
    for index, row in X.iterrows():
        prediction, terms = _run_procedure_queries(
            fixture.connection, {**row.to_dict(), "Exposure": exposure[index]}
        )
        actual.append(prediction)
        assert len(terms) == 4
        assert {term[1] for term in terms} == {"SPLINE", "NUMERIC", "CELL", "OFFSET_PER_UNIT"}
    np.testing.assert_allclose(actual, expected, rtol=1e-10, atol=1e-12)


@pytest.mark.parametrize("marker", [None, "UNKNOWN"])
def test_ambiguous_or_invalid_offset_marker_rejected_before_publication(
    native_package, tmp_path, marker
):
    fixture = native_package
    path = tmp_path / "native.xlsx"
    workbook = load_workbook(path)
    sheet = workbook["Rating Tables"]
    column = next(c for c in range(1, sheet.max_column + 1) if sheet.cell(5, c).value == "Exposure")
    sheet.cell(7, column).comment = (
        None
        if marker is None
        else Comment(f"pricing_pipeline:offset_representation={marker}", "pytest")
    )
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="ambiguous per_unit|unsupported offset representation"):
        _verified_rating_frames(
            workbook_path=path,
            export_id="invalid",
            model_name="NATIVE",
            model_version="v1",
            effective_from=None,
            publication_receipt_path=fixture.receipt_path,
            publication_receipt_sha256=fixture.receipt_sha256,
        )


@pytest.mark.parametrize("value", [None, "bad", "per_unit", 0, -1])
def test_per_unit_offset_rejects_invalid_input_without_default(native_package, value):
    fixture = native_package
    # Prove defaults cannot hide an invalid continuous offset input.
    fixture.connection.execute(
        text(
            "UPDATE pricing.PRICING_COMPILED_RATE_CELL SET is_default = 1 WHERE term_type = 'OFFSET_FACTOR'"
        )
    )
    features = {**fixture.X.iloc[0].to_dict(), "Exposure": value}
    with pytest.raises(ValueError, match="did not match every required term"):
        _run_procedure_queries(fixture.connection, features)


@pytest.mark.parametrize("feature", ["Exposure", "density", "numeric", "region"])
def test_native_package_requires_every_input(native_package, feature):
    fixture = native_package
    features = {**fixture.X.iloc[0].to_dict(), "Exposure": 0.5}
    del features[feature]
    with pytest.raises(ValueError, match="did not match every required term"):
        _run_procedure_queries(fixture.connection, features)


def test_error_boundary_rejects_uncovered_spline_inputs(native_package):
    fixture = native_package
    if fixture.boundary != "error":
        return
    for value in (-0.1, 5.0):
        X = fixture.X.iloc[[0]].copy()
        X["density"] = value
        with pytest.raises(ValueError):
            fixture.model.predict(X, offset=np.log([0.5 * fixture.scale]))
        with pytest.raises(ValueError, match="did not match every required term"):
            _run_procedure_queries(fixture.connection, {**X.iloc[0].to_dict(), "Exposure": 0.5})


def test_offset_breakdown_reports_evaluated_multiplier(native_package):
    fixture = native_package
    _run_procedure_queries(fixture.connection, {**fixture.X.iloc[0].to_dict(), "Exposure": 0.5})
    sql = _active_procedure()
    breakdown = re.search(r"IF @include_breakdown = 1\s+BEGIN\s+(SELECT.*?);", sql, re.DOTALL)[1]
    rows = fixture.connection.connection.driver_connection.execute(
        _sqlite_query(breakdown)
    ).fetchall()
    offset = next(row for row in rows if row[3] == "OFFSET_PER_UNIT")
    assert offset[-2] == pytest.approx(0.5 * fixture.scale, rel=1e-12)
    assert offset[-1] == pytest.approx(math.log(0.5 * fixture.scale), rel=1e-12)


def test_discrete_offset_lookup_keeps_native_predictions(native_package, tmp_path):
    fixture = native_package
    path = export_rating_tables(
        fixture.model,
        fixture.X,
        fixture.y,
        None,
        tmp_path / "discrete.xlsx",
        offset=np.log(fixture.exposure * fixture.scale),
        offset_source=fixture.exposure,
        offset_name="Exposure",
        offset_kind="discrete",
        offset_max_exact_levels=300,
    )
    frames = _verified_rating_frames(
        workbook_path=path,
        export_id="discrete",
        model_name="NATIVE",
        model_version="v2",
        effective_from=None,
        publication_receipt_path=fixture.receipt_path,
        publication_receipt_sha256=fixture.receipt_sha256,
    )
    _insert_frames_package(fixture.connection, frames, 2)
    value = frames[2].loc[frames[2].feature_name.eq("Exposure"), "level_code"].iloc[0]
    X = fixture.X.iloc[[7]]
    actual, terms = _run_procedure_queries(
        fixture.connection, {**X.iloc[0].to_dict(), "Exposure": value}, package=2
    )
    expected = fixture.model.predict(X, offset=np.log([float(value) * fixture.scale]))[0]
    assert actual == pytest.approx(expected, rel=1e-10, abs=1e-12)
    assert ("Exposure", "CELL") in {(term[0], term[1]) for term in terms}
    # A second package with a different offset representation cannot change the first.
    original, terms = _run_procedure_queries(
        fixture.connection, {**X.iloc[0].to_dict(), "Exposure": float(value)}, package=1
    )
    assert original == pytest.approx(expected, rel=1e-10, abs=1e-12)
    assert ("Exposure", "OFFSET_PER_UNIT") in {(term[0], term[1]) for term in terms}


def test_persisted_transform_preparation_matches_native_prediction(native_package):
    fixture = native_package
    metadata = fixture.connection.execute(
        text("SELECT transforms_json FROM pricing.V_MODEL_SPLINE_SEGMENT LIMIT 1")
    ).scalar_one()
    prepared = apply_transforms(
        fixture.raw_frame.iloc[[5, 50, 110]], transforms_from_metadata(json.loads(metadata))
    )
    expected = fixture.model.predict(
        prepared, offset=np.log(np.full(len(prepared), 0.5 * fixture.scale))
    )
    actual = [
        _run_procedure_queries(fixture.connection, {**row.to_dict(), "Exposure": 0.5})[0]
        for _, row in prepared.iterrows()
    ]
    np.testing.assert_allclose(actual, expected, rtol=1e-10, atol=1e-12)


@pytest.mark.parametrize("labels", [["per_unit", "annual"], ["per_unit"]])
def test_discrete_offset_label_per_unit_is_not_a_numeric_factor(native_package, tmp_path, labels):
    fixture = native_package
    source = pd.Series(np.resize(labels, len(fixture.X)), name="Exposure")
    factors = {label: 0.2 * (index + 1) for index, label in enumerate(labels)}
    offset = np.log(source.map(factors).to_numpy())
    path = export_rating_tables(
        fixture.model,
        fixture.X,
        fixture.y,
        None,
        tmp_path / "labelled-offset.xlsx",
        offset=offset,
        offset_source=source,
        offset_name="Exposure",
        offset_kind="discrete",
    )
    frames = _verified_rating_frames(
        workbook_path=path,
        export_id="labelled",
        model_name="NATIVE",
        model_version="v2",
        effective_from=None,
        publication_receipt_path=fixture.receipt_path,
        publication_receipt_sha256=fixture.receipt_sha256,
    )
    _insert_frames_package(fixture.connection, frames, 2)
    metadata = fixture.connection.execute(
        text(
            "SELECT term_metadata_json FROM pricing.PRICING_TERM "
            "WHERE rate_package_id = 2 AND term_type = 'OFFSET_FACTOR'"
        )
    ).scalar_one()
    assert json.loads(metadata)["rating_representation"] == "LOOKUP"
    X = fixture.X.iloc[[7]]
    with pytest.raises(ValueError, match="did not match every required term"):
        _run_procedure_queries(
            fixture.connection, {**X.iloc[0].to_dict(), "Exposure": 4}, package=2
        )
    for label in labels:
        actual, terms = _run_procedure_queries(
            fixture.connection,
            {**X.iloc[0].to_dict(), "Exposure": label},
            package=2,
        )
        expected = fixture.model.predict(X, offset=np.log([factors[label]]))[0]
        assert actual == pytest.approx(expected, rel=1e-10, abs=1e-12)
        assert ("Exposure", "CELL") in {(term[0], term[1]) for term in terms}


def test_legacy_unmarked_offset_keeps_lookup_semantics(native_package):
    fixture = native_package
    fixture.connection.execute(
        text(
            "UPDATE pricing.PRICING_TERM SET term_metadata_json = NULL WHERE term_type = 'OFFSET_FACTOR'"
        )
    )
    X = fixture.X.iloc[[7]]
    with pytest.raises(ValueError, match="did not match every required term"):
        _run_procedure_queries(fixture.connection, {**X.iloc[0].to_dict(), "Exposure": 4})
    actual, terms = _run_procedure_queries(
        fixture.connection, {**X.iloc[0].to_dict(), "Exposure": "per_unit"}
    )
    expected = fixture.model.predict(X, offset=np.log([fixture.scale]))[0]
    assert actual == pytest.approx(expected, rel=1e-10, abs=1e-12)
    assert ("Exposure", "CELL") in {(term[0], term[1]) for term in terms}
